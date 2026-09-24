"""DOCX parsing via ``python-docx``. Headings come from paragraph style
names ("Heading 1".."Heading 6") — reliable, since Word actually stores
that. No page markers: DOCX doesn't store real page boundaries (pagination
is computed at render/print time), so citations for file-type ``docx``
sources fall back to char ranges, same as HTML/Markdown/text.
"""

from __future__ import annotations

import io
import re

from docx import Document as DocxDocument

from app.rag.parsers.base import Heading, ParsedDocument

_HEADING_STYLE = re.compile(r"^Heading (\d)$")


def parse_docx(content: bytes) -> ParsedDocument:
    doc = DocxDocument(io.BytesIO(content))
    parts: list[str] = []
    headings: list[Heading] = []
    offset = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        m = _HEADING_STYLE.match(para.style.name if para.style else "")
        if m:
            headings.append(Heading(level=int(m.group(1)), text=text, offset=offset))
        parts.append(text)
        offset += len(text) + 2

    title = headings[0].text if headings and headings[0].level == 1 else None
    return ParsedDocument(text="\n\n".join(parts), title=title, headings=headings)


__all__ = ["parse_docx"]
